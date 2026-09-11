#!/usr/bin/env python3
"""Build Northline_Unit_Economics.xlsx — SKU/channel contribution, ramp, and payback portfolio sample."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.utils import get_column_letter

yellow = PatternFill("solid", fgColor="FFF2CC")
header_fill = PatternFill("solid", fgColor="1F4E79")
section_fill = PatternFill("solid", fgColor="D6E3F0")
green_fill = PatternFill("solid", fgColor="C6EFCE")
amber_fill = PatternFill("solid", fgColor="FFE699")
red_fill = PatternFill("solid", fgColor="F8CBAD")
tile_fill = PatternFill("solid", fgColor="E9EDF4")
dark_navy = PatternFill("solid", fgColor="0D2B4A")
accent_fill = PatternFill("solid", fgColor="2E75B6")

input_font = Font(name="Calibri", size=11, color="0000FF")
black = Font(name="Calibri", size=11, color="000000")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
title_font = Font(name="Calibri", size=16, bold=True, color="1F4E79")
section_font = Font(name="Calibri", size=12, bold=True, color="1F4E79")
bold = Font(name="Calibri", size=11, bold=True)
bold_black = Font(name="Calibri", size=11, bold=True, color="000000")
italic_grey = Font(name="Calibri", size=10, italic=True, color="666666")
small_grey = Font(name="Calibri", size=9, italic=True, color="666666")
tile_title = Font(name="Calibri", size=9, bold=True, color="1F4E79")
tile_value = Font(name="Calibri", size=14, bold=True, color="000000")
white_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
big_title = Font(name="Calibri", size=18, bold=True, color="FFFFFF")
subtitle_w = Font(name="Calibri", size=11, color="D6E3F0")

thin = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)

money = '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)'
money0 = '_($* #,##0_);_($* (#,##0);_($* "-"??_);_(@_)'
pct = "0.0%"
num = "#,##0"
num1 = "#,##0.0"

COVER = "00_Cover"
ASSUMP = "01_Assumptions"
UNIT = "02_Unit_P&L"
CHAN = "03_Channel_View"
RAMP = "04_Cohort_or_Ramp"
DASH = "05_Dashboard"
DICT = "06_Data_Dictionary"

# SKUs — fictional Northline CPG
# price, cogs, var_selling (per unit $), slotting_or_cac (one-time $000s or per-unit where noted),
# annual_volume (000 units), channel mix defaults handled on channel tab
SKUS = [
    # code, name, price, cogs, var_sell, slotting_$000, annual_vol_000, launch?
    ("NL-CRUNCH", "Northline Crunch Bar", 2.49, 0.92, 0.18, 85, 4200, False),
    ("NL-PROTEIN", "Protein Trail Mix", 4.99, 1.85, 0.35, 120, 2100, False),
    ("NL-OAT", "Overnight Oats Cup", 3.29, 1.15, 0.22, 95, 3100, False),
    ("NL-KIDS", "Kids Fruit Bites", 2.99, 1.05, 0.28, 150, 2800, False),
    ("NL-PROBIO", "Probiotic Bar (launch)", 3.79, 1.42, 0.40, 220, 900, True),
    ("NL-COLD", "Cold Brew RTD", 2.79, 0.88, 0.25, 110, 3500, False),
]

CHANNELS = ["Grocery", "Club", "eComm", "Convenience"]
# Channel mix % by SKU (rows align with SKUS) — must sum ~1.0
CHANNEL_MIX = [
    [0.45, 0.25, 0.15, 0.15],  # CRUNCH
    [0.35, 0.35, 0.20, 0.10],  # PROTEIN
    [0.50, 0.15, 0.25, 0.10],  # OAT
    [0.40, 0.30, 0.10, 0.20],  # KIDS
    [0.30, 0.20, 0.40, 0.10],  # PROBIO launch — eComm heavy
    [0.25, 0.20, 0.15, 0.40],  # COLD — convenience heavy
]
# Channel-specific variable take rates (additive to base var_sell as % of price) — e.g. retail margin / fees
CHANNEL_FEE_PCT = [0.28, 0.22, 0.18, 0.32]  # grocery higher slotting already separate; fee = retailer margin proxy

# 12-month ramp for launch SKU (NL-PROBIO) — monthly volume 000 units
RAMP_VOL = [20, 35, 55, 70, 80, 90, 95, 100, 105, 110, 115, 125]  # sum = 1000 ≈ annual 900 + ramp
# Fixed launch spend $000s by month (slotting front-loaded + CAC)
RAMP_FIXED = [120, 40, 25, 15, 10, 5, 5, 0, 0, 0, 0, 0]


def style_input(cell):
    cell.fill = yellow
    cell.font = input_font
    cell.border = thin
    cell.alignment = Alignment(horizontal="center")


def style_formula(cell, key=False):
    cell.font = bold_black if key else black
    cell.border = thin
    cell.alignment = Alignment(horizontal="center")
    if key:
        cell.fill = green_fill


def style_header_cell(cell, value):
    cell.value = value
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
    cell.border = thin


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def landscape(ws, fit_height=1):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = fit_height
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.sheet_view.showGridLines = False
    ws.page_setup.horizontalCentered = True
    ws.oddFooter.left.text = "Northline Consumer Products  |  fictional sample  |  $000s where noted"
    ws.oddFooter.right.text = "CONFIDENTIAL — portfolio sample"


wb = Workbook()

# ═══════════════════════════════════════════════════════════════════
# 00_Cover
# ═══════════════════════════════════════════════════════════════════
ws = wb.active
ws.title = COVER
landscape(ws)
set_col_widths(ws, [3, 24, 55, 18, 18])

ws.merge_cells("B2:E2")
ws["B2"] = "NORTHLINE CONSUMER PRODUCTS"
ws["B2"].font = title_font
ws.merge_cells("B3:E3")
ws["B3"] = "Unit Economics Workbook — SKU contribution, channel view, launch ramp"
ws["B3"].font = Font(name="Calibri", size=14, bold=True, color="2E75B6")
ws.merge_cells("B4:E4")
ws["B4"] = "Fictional CPG portfolio sample  ·  per-unit $ and $000s  ·  formulas only (no VBA)"
ws["B4"].font = italic_grey

ws["B6"] = "Purpose"
ws["B6"].font = section_font
ws.merge_cells("B7:E7")
ws["B7"] = (
    "Per-SKU unit P&L (price → contribution), channel mix and contribution, "
    "and a 12-month ramp / payback on the launch SKU (NL-Probiotic Bar). "
    "Built so FP&A can stress price, COGS, and CAC before a category review."
)
ws["B7"].alignment = Alignment(wrap_text=True)
ws.row_dimensions[7].height = 42

ws["B9"] = "How to use"
ws["B9"].font = section_font
steps = [
    "1. Open 01_Assumptions — edit SKU price, COGS, variable selling, slotting, and annual volume (yellow).",
    "2. Read 02_Unit_P&L — per-unit contribution and contribution margin % by SKU (formulas).",
    "3. Review 03_Channel_View — mix and channel contribution; fees are yellow inputs.",
    "4. Stress 04_Cohort_or_Ramp — 12-month volume ramp and cumulative payback on the launch SKU.",
    "5. Use 05_Dashboard as the one-pager for a category or innovation review.",
]
for i, s in enumerate(steps):
    ws.cell(10 + i, 2, s).font = black
    ws.merge_cells(start_row=10 + i, start_column=2, end_row=10 + i, end_column=5)

ws["B16"] = "Color convention"
ws["B16"].font = section_font
ws["B17"] = "Yellow fill + blue font"
style_input(ws["B17"])
ws["C17"] = "= inputs (edit these)"
ws["C17"].font = black
ws["B18"] = "Black font"
style_formula(ws["B18"])
ws["C18"] = "= formulas (do not overwrite)"
ws["C18"].font = black
ws["B19"] = "Green highlight"
ws["B19"].fill = green_fill
ws["B19"].border = thin
ws["C19"] = "= key outputs (contribution, payback)"
ws["C19"].font = black

ws["B21"] = "Tabs"
ws["B21"].font = section_font
tabs = [
    (COVER, "Purpose and navigation"),
    (ASSUMP, "SKU list, price, COGS, var selling, slotting, volume"),
    (UNIT, "Per-unit contribution by SKU"),
    (CHAN, "Channel mix and contribution"),
    (RAMP, "12-month launch ramp / payback"),
    (DASH, "Dashboard one-pager"),
    (DICT, "Field definitions"),
]
for i, (t, d) in enumerate(tabs):
    r = 22 + i
    ws.cell(r, 2, t).font = bold
    ws.cell(r, 2).border = thin
    ws.cell(r, 3, d).border = thin
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)

ws["B30"] = "Sai Siri Bandaru — Financial Analyst | FP&A | forecasting, variance analysis, Excel"
ws["B30"].font = small_grey
ws["B31"] = "All sample numbers are fictional. No employer data."
ws["B31"].font = small_grey

# ═══════════════════════════════════════════════════════════════════
# 01_Assumptions
# ═══════════════════════════════════════════════════════════════════
ws = wb.create_sheet(ASSUMP)
landscape(ws)
set_col_widths(ws, [3, 14, 26, 10, 10, 12, 12, 14, 10, 12, 36])

ws["B2"] = "01  ·  Assumptions"
ws["B2"].font = title_font
ws["B3"] = "SKU economics inputs. Yellow = edit. Downstream tabs are formulas."
ws["B3"].font = italic_grey

ws["B5"] = "Company"
ws["C5"] = "Northline Consumer Products"
style_input(ws["C5"])
ws.merge_cells("C5:E5")
ws["B6"] = "Fiscal year"
ws["C6"] = "FY26"
style_input(ws["C6"])
ws["B7"] = "Currency / units"
ws["C7"] = "USD; volume in 000 units; slotting in $000s"
style_input(ws["C7"])
ws.merge_cells("C7:F7")

headers = ["SKU", "Product", "Price $", "COGS $", "Var sell $", "Slotting $000",
           "Ann vol (000)", "Launch?", "WACC / hurdle", "Notes"]
for i, h in enumerate(headers):
    style_header_cell(ws.cell(9, 2 + i), h)
ws.row_dimensions[9].height = 32

for i, (code, name, price, cogs, vsell, slot, vol, launch) in enumerate(SKUS):
    r = 10 + i
    c = ws.cell(r, 2, code)
    style_input(c)
    c = ws.cell(r, 3, name)
    style_input(c)
    c.alignment = Alignment(horizontal="left")
    c = ws.cell(r, 4, price)
    style_input(c)
    c.number_format = money
    c = ws.cell(r, 5, cogs)
    style_input(c)
    c.number_format = money
    c = ws.cell(r, 6, vsell)
    style_input(c)
    c.number_format = money
    c = ws.cell(r, 7, slot)
    style_input(c)
    c.number_format = money0
    c = ws.cell(r, 8, vol)
    style_input(c)
    c.number_format = num
    c = ws.cell(r, 9, "Y" if launch else "N")
    style_input(c)
    if i == 0:
        c = ws.cell(r, 10, 0.10)
        style_input(c)
        c.number_format = pct
    else:
        c = ws.cell(r, 10, f"=J10")
        style_formula(c)
        c.number_format = pct
    note = "Launch SKU — see ramp tab" if launch else "Core assortment"
    ws.cell(r, 11, note).font = small_grey
    ws.cell(r, 11).border = thin

ws["B17"] = "Global toggles"
ws["B17"].font = section_font
ws["B17"].fill = section_fill
ws.merge_cells("B17:D17")

ws["B18"] = "Price stress %"
ws["C18"] = 0.0
style_input(ws["C18"])
ws["C18"].number_format = pct
ws["D18"] = "Applied on Unit P&L as Price × (1 + stress)"
ws["D18"].font = small_grey

ws["B19"] = "COGS stress %"
ws["C19"] = 0.0
style_input(ws["C19"])
ws["C19"].number_format = pct
ws["D19"] = "Applied on Unit P&L as COGS × (1 + stress)"
ws["D19"].font = small_grey

ws["B20"] = "Target contribution margin %"
ws["C20"] = 0.35
style_input(ws["C20"])
ws["C20"].number_format = pct
ws["D20"] = "Used for RAG on Unit P&L and Dashboard"
ws["D20"].font = small_grey

ws["B21"] = "Payback hurdle (months)"
ws["C21"] = 12
style_input(ws["C21"])
ws["D21"] = "Launch SKU cumulative contribution vs slotting+CAC"
ws["D21"].font = small_grey

ws["B23"] = "Per-unit $ are shelf / list before channel fees. Slotting is annualized $000s (one-time for launch on ramp tab)."
ws["B23"].font = small_grey

# ═══════════════════════════════════════════════════════════════════
# 02_Unit_P&L
# ═══════════════════════════════════════════════════════════════════
ws = wb.create_sheet(UNIT)
landscape(ws)
set_col_widths(ws, [3, 14, 22, 11, 11, 11, 12, 12, 12, 14, 14, 12, 14])

ws["B2"] = "02  ·  Unit P&L"
ws["B2"].font = title_font
ws["B3"] = "Per-unit bridge: stressed price → COGS → var sell → contribution. Annual $000s = unit × volume."
ws["B3"].font = italic_grey

headers = ["SKU", "Product", "Price $", "COGS $", "Var sell $",
           "Unit contrib $", "CM %", "vs Target", "RAG",
           "Ann vol (000)", "Ann contrib $000", "Slotting $000", "Ann CM after slot $000"]
for i, h in enumerate(headers):
    style_header_cell(ws.cell(5, 2 + i), h)
ws.row_dimensions[5].height = 36

for i in range(len(SKUS)):
    r = 6 + i
    a = 10 + i  # assumptions row
    # SKU / Product
    ws.cell(r, 2, f"='01_Assumptions'!B{a}")
    style_formula(ws.cell(r, 2))
    ws.cell(r, 3, f"='01_Assumptions'!C{a}")
    style_formula(ws.cell(r, 3))
    ws.cell(r, 3).alignment = Alignment(horizontal="left")

    # Stressed price
    c = ws.cell(r, 4, f"='01_Assumptions'!D{a}*(1+'01_Assumptions'!C18)")
    style_formula(c)
    c.number_format = money
    # Stressed COGS
    c = ws.cell(r, 5, f"='01_Assumptions'!E{a}*(1+'01_Assumptions'!C19)")
    style_formula(c)
    c.number_format = money
    # Var sell
    c = ws.cell(r, 6, f"='01_Assumptions'!F{a}")
    style_formula(c)
    c.number_format = money
    # Unit contribution
    c = ws.cell(r, 7, f"=D{r}-E{r}-F{r}")
    style_formula(c, key=True)
    c.number_format = money
    # CM %
    c = ws.cell(r, 8, f"=IF(D{r}=0,0,G{r}/D{r})")
    style_formula(c, key=True)
    c.number_format = pct
    # vs Target (pp)
    c = ws.cell(r, 9, f"=H{r}-'01_Assumptions'!C20")
    style_formula(c)
    c.number_format = pct
    # RAG
    c = ws.cell(r, 10,
                f'=IF(H{r}>=\'01_Assumptions\'!C20,"Green",'
                f'IF(H{r}>=\'01_Assumptions\'!C20-0.05,"Amber","Red"))')
    style_formula(c, key=True)
    # Volume
    c = ws.cell(r, 11, f"='01_Assumptions'!H{a}")
    style_formula(c)
    c.number_format = num
    # Ann contrib $000 = unit contrib $ × vol_000
    c = ws.cell(r, 12, f"=G{r}*K{r}")
    style_formula(c, key=True)
    c.number_format = money0
    # Slotting
    c = ws.cell(r, 13, f"='01_Assumptions'!G{a}")
    style_formula(c)
    c.number_format = money0
    # After slotting
    c = ws.cell(r, 14, f"=L{r}-M{r}")
    style_formula(c, key=True)
    c.number_format = money0

ws.conditional_formatting.add("J6:J11", FormulaRule(formula=['J6="Green"'], fill=green_fill))
ws.conditional_formatting.add("J6:J11", FormulaRule(formula=['J6="Amber"'], fill=amber_fill))
ws.conditional_formatting.add("J6:J11", FormulaRule(formula=['J6="Red"'], fill=red_fill))

# Totals
ws["B13"] = "TOTAL / avg"
ws["B13"].font = bold_black
ws["B13"].border = thin
ws["K13"] = "=SUM(K6:K11)"
style_formula(ws["K13"], key=True)
ws["K13"].number_format = num
ws["L13"] = "=SUM(L6:L11)"
style_formula(ws["L13"], key=True)
ws["L13"].number_format = money0
ws["M13"] = "=SUM(M6:M11)"
style_formula(ws["M13"])
ws["M13"].number_format = money0
ws["N13"] = "=SUM(N6:N11)"
style_formula(ws["N13"], key=True)
ws["N13"].number_format = money0
ws["H13"] = "=IF(SUMPRODUCT(D6:D11,K6:K11)=0,0,SUMPRODUCT(G6:G11,K6:K11)/SUMPRODUCT(D6:D11,K6:K11))"
style_formula(ws["H13"], key=True)
ws["H13"].number_format = pct

ws["B15"] = "Unit contrib $ = Price − COGS − Var sell (after stress toggles). Ann contrib $000 = unit $ × volume (000)."
ws["B15"].font = small_grey

# ═══════════════════════════════════════════════════════════════════
# 03_Channel_View
# ═══════════════════════════════════════════════════════════════════
ws = wb.create_sheet(CHAN)
landscape(ws)
set_col_widths(ws, [3, 14, 12, 12, 12, 12, 12, 14, 14, 14, 14, 14])

ws["B2"] = "03  ·  Channel View"
ws["B2"].font = title_font
ws["B3"] = "Mix % by SKU (yellow) and channel fee % of price (yellow). Contribution after channel fee."
ws["B3"].font = italic_grey

ws["B5"] = "Channel fee % of price (retailer / platform take)"
ws["B5"].font = section_font
ws["B5"].fill = section_fill
ws.merge_cells("B5:F5")

style_header_cell(ws.cell(6, 2), "Metric")
for i, ch in enumerate(CHANNELS):
    style_header_cell(ws.cell(6, 3 + i), ch)

ws["B7"] = "Fee % of price"
ws["B7"].border = thin
for i, fee in enumerate(CHANNEL_FEE_PCT):
    c = ws.cell(7, 3 + i, fee)
    style_input(c)
    c.number_format = pct

ws["B9"] = "Channel mix by SKU (must sum to ~100%)"
ws["B9"].font = section_font
ws["B9"].fill = section_fill
ws.merge_cells("B9:G9")

headers = ["SKU"] + CHANNELS + ["Mix check", "Unit CM $ (pre-fee)", "Blended fee %", "Unit CM $ after fee", "Ann CM $000"]
for i, h in enumerate(headers):
    style_header_cell(ws.cell(10, 2 + i), h)
ws.row_dimensions[10].height = 32

for i in range(len(SKUS)):
    r = 11 + i
    a = 10 + i
    ws.cell(r, 2, f"='01_Assumptions'!B{a}")
    style_formula(ws.cell(r, 2))
    for j, mix in enumerate(CHANNEL_MIX[i]):
        c = ws.cell(r, 3 + j, mix)
        style_input(c)
        c.number_format = pct
    # Mix check
    c = ws.cell(r, 7, f"=C{r}+D{r}+E{r}+F{r}")
    style_formula(c)
    c.number_format = pct
    # Unit CM pre-fee from Unit P&L
    c = ws.cell(r, 8, f"='02_Unit_P&L'!G{6+i}")
    style_formula(c)
    c.number_format = money
    # Blended fee %
    c = ws.cell(r, 9, f"=C{r}*$C$7+D{r}*$D$7+E{r}*$E$7+F{r}*$F$7")
    style_formula(c)
    c.number_format = pct
    # Unit CM after fee = Price*(1) path: unit contrib already net of var_sell;
    # channel fee is additional % of price
    c = ws.cell(r, 10, f"=H{r}-'02_Unit_P&L'!D{6+i}*I{r}")
    style_formula(c, key=True)
    c.number_format = money
    # Ann CM $000
    c = ws.cell(r, 11, f"=J{r}*'02_Unit_P&L'!K{6+i}")
    style_formula(c, key=True)
    c.number_format = money0

ws["B18"] = "Channel contribution $000s (volume × mix × unit CM after fee)"
ws["B18"].font = section_font
ws["B18"].fill = section_fill
ws.merge_cells("B18:F18")

style_header_cell(ws.cell(19, 2), "SKU")
for i, ch in enumerate(CHANNELS):
    style_header_cell(ws.cell(19, 3 + i), ch)
style_header_cell(ws.cell(19, 7), "Total")

for i in range(len(SKUS)):
    r = 20 + i
    src = 11 + i
    ws.cell(r, 2, f"=B{src}")
    style_formula(ws.cell(r, 2))
    for j in range(4):
        # vol * mix * unit CM after fee
        mix_col = get_column_letter(3 + j)
        c = ws.cell(r, 3 + j, f"='02_Unit_P&L'!K{6+i}*{mix_col}{src}*$J{src}")
        style_formula(c)
        c.number_format = money0
    c = ws.cell(r, 7, f"=C{r}+D{r}+E{r}+F{r}")
    style_formula(c, key=True)
    c.number_format = money0

ws["B27"] = "TOTAL"
ws["B27"].font = bold_black
ws["B27"].border = thin
for j in range(4):
    col = get_column_letter(3 + j)
    c = ws.cell(27, 3 + j, f"=SUM({col}20:{col}25)")
    style_formula(c, key=True)
    c.number_format = money0
c = ws.cell(27, 7, "=SUM(G20:G25)")
style_formula(c, key=True)
c.number_format = money0

ws["B29"] = "Mix check should be ~100%. If not, channel $000s are still proportional but flag for cleanup."
ws["B29"].font = small_grey

# ═══════════════════════════════════════════════════════════════════
# 04_Cohort_or_Ramp
# ═══════════════════════════════════════════════════════════════════
ws = wb.create_sheet(RAMP)
landscape(ws)
set_col_widths(ws, [3, 22] + [10] * 12 + [12, 12])

ws["B2"] = "04  ·  Cohort / Ramp — NL-PROBIO launch"
ws["B2"].font = title_font
ws["B3"] = "12-month volume ramp and cumulative contribution vs launch spend (slotting + CAC). Payback month highlighted."
ws["B3"].font = italic_grey

ws["B5"] = "Launch SKU"
ws["C5"] = "NL-PROBIO"
style_input(ws["C5"])
ws["B6"] = "Unit price $ (linked)"
ws["C6"] = "='02_Unit_P&L'!D10"  # NL-PROBIO is 5th SKU → row 10 on unit (6+4)
style_formula(ws["C6"])
ws["C6"].number_format = money
ws["B7"] = "Unit contrib $ (pre-fee)"
ws["C7"] = "='02_Unit_P&L'!G10"
style_formula(ws["C7"], key=True)
ws["C7"].number_format = money
ws["B8"] = "Blended fee % (from channel)"
ws["C8"] = "='03_Channel_View'!I15"  # row 11+4=15
style_formula(ws["C8"])
ws["C8"].number_format = pct
ws["B9"] = "Unit CM $ after fee"
ws["C9"] = "=C7-C6*C8"
style_formula(ws["C9"], key=True)
ws["C9"].number_format = money
ws["B10"] = "Payback hurdle (mo)"
ws["C10"] = "='01_Assumptions'!C21"
style_formula(ws["C10"])

months = [f"M{i}" for i in range(1, 13)]
style_header_cell(ws.cell(12, 2), "Metric")
for i, m in enumerate(months):
    style_header_cell(ws.cell(12, 3 + i), m)
style_header_cell(ws.cell(12, 15), "FY Total")

# Volume row (inputs)
ws["B13"] = "Volume (000 units)"
ws["B13"].border = thin
for i, v in enumerate(RAMP_VOL):
    c = ws.cell(13, 3 + i, v)
    style_input(c)
    c.number_format = num1
ws["O13"] = "=SUM(C13:N13)"
style_formula(ws["O13"], key=True)
ws["O13"].number_format = num1

# Launch spend $000s (inputs)
ws["B14"] = "Launch spend $000s"
ws["B14"].border = thin
for i, v in enumerate(RAMP_FIXED):
    c = ws.cell(14, 3 + i, v)
    style_input(c)
    c.number_format = money0
ws["O14"] = "=SUM(C14:N14)"
style_formula(ws["O14"])
ws["O14"].number_format = money0

# Revenue $000s = price * vol
ws["B15"] = "Revenue $000s"
ws["B15"].border = thin
for i in range(12):
    col = get_column_letter(3 + i)
    c = ws.cell(15, 3 + i, f"=$C$6*{col}13")
    style_formula(c)
    c.number_format = money0
ws["O15"] = "=SUM(C15:N15)"
style_formula(ws["O15"])
ws["O15"].number_format = money0

# Contribution $000s = unit CM after fee * vol
ws["B16"] = "Contribution $000s"
ws["B16"].border = thin
for i in range(12):
    col = get_column_letter(3 + i)
    c = ws.cell(16, 3 + i, f"=$C$9*{col}13")
    style_formula(c, key=True)
    c.number_format = money0
ws["O16"] = "=SUM(C16:N16)"
style_formula(ws["O16"], key=True)
ws["O16"].number_format = money0

# Net of launch spend
ws["B17"] = "Net of launch spend $000s"
ws["B17"].border = thin
for i in range(12):
    col = get_column_letter(3 + i)
    c = ws.cell(17, 3 + i, f"={col}16-{col}14")
    style_formula(c)
    c.number_format = money0
ws["O17"] = "=SUM(C17:N17)"
style_formula(ws["O17"], key=True)
ws["O17"].number_format = money0

# Cumulative net
ws["B18"] = "Cumulative net $000s"
ws["B18"].border = thin
ws["C18"] = "=C17"
style_formula(ws["C18"], key=True)
ws["C18"].number_format = money0
for i in range(1, 12):
    col = get_column_letter(3 + i)
    prev = get_column_letter(2 + i)
    c = ws.cell(18, 3 + i, f"={prev}18+{col}17")
    style_formula(c, key=True)
    c.number_format = money0
ws["O18"] = "=N18"
style_formula(ws["O18"], key=True)
ws["O18"].number_format = money0

# Payback flag
ws["B19"] = "Payback reached?"
ws["B19"].border = thin
for i in range(12):
    col = get_column_letter(3 + i)
    if i == 0:
        c = ws.cell(19, 3 + i, f'=IF({col}18>=0,"YES","")')
    else:
        prev = get_column_letter(2 + i)
        c = ws.cell(19, 3 + i, f'=IF(AND({prev}18<0,{col}18>=0),"YES",IF({col}18>=0,"",""))')
    style_formula(c, key=True)

ws.conditional_formatting.add(
    "C19:N19",
    FormulaRule(formula=['C19="YES"'], fill=green_fill),
)
ws.conditional_formatting.add(
    "C18:N18",
    CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=green_fill),
)
ws.conditional_formatting.add(
    "C18:N18",
    CellIsRule(operator="lessThan", formula=["0"], fill=red_fill),
)

ws["B21"] = "Payback month #"
ws["C21"] = '=IFERROR(MATCH("YES",C19:N19,0),"Not within 12 mo")'
style_formula(ws["C21"], key=True)
ws["B22"] = "Within hurdle?"
ws["C22"] = '=IF(ISNUMBER(C21),IF(C21<=C10,"Yes — inside hurdle","No — beyond hurdle"),"No — not paid back")'
style_formula(ws["C22"], key=True)
ws["B23"] = "Year-1 CM after launch spend $000s"
ws["C23"] = "=O17"
style_formula(ws["C23"], key=True)
ws["C23"].number_format = money0

ws["B25"] = "Edit yellow volume and launch spend to stress payback. Unit CM flows from Assumptions → Unit P&L → Channel fees."
ws["B25"].font = small_grey

# Chart data is the rows themselves
chart = LineChart()
chart.title = "Cumulative net contribution ($000s)"
chart.style = 10
chart.y_axis.title = "$000s"
chart.x_axis.title = "Month"
data = Reference(ws, min_col=2, min_row=18, max_col=14, max_row=18)
cats = Reference(ws, min_col=3, min_row=12, max_col=14)
chart.add_data(data, from_rows=True, titles_from_data=True)
chart.set_categories(cats)
chart.width = 18
chart.height = 8
ws.add_chart(chart, "B27")

bar = BarChart()
bar.type = "col"
bar.title = "Monthly contribution vs launch spend"
bar.style = 10
data2 = Reference(ws, min_col=2, min_row=14, max_col=14, max_row=16)
bar.add_data(data2, from_rows=True, titles_from_data=True)
bar.set_categories(cats)
bar.width = 18
bar.height = 8
ws.add_chart(bar, "B42")

# ═══════════════════════════════════════════════════════════════════
# 05_Dashboard
# ═══════════════════════════════════════════════════════════════════
ws = wb.create_sheet(DASH)
landscape(ws, fit_height=1)
set_col_widths(ws, [2.5, 14, 12, 12, 12, 12, 12, 3, 14, 12, 12, 12, 12, 2.5])

for col in range(2, 14):
    ws.cell(1, col).fill = dark_navy
ws.merge_cells("B1:M1")
ws["B1"] = '="NORTHLINE  ·  UNIT ECONOMICS DASHBOARD  ·  "&\'01_Assumptions\'!C6'
ws["B1"].font = big_title
ws["B1"].fill = dark_navy
ws.row_dimensions[1].height = 28

for col in range(2, 14):
    ws.cell(2, col).fill = dark_navy
ws.merge_cells("B2:M2")
ws["B2"] = "='01_Assumptions'!C5&\"  ·  per-unit $ and $000s  ·  fictional sample\""
ws["B2"].font = subtitle_w
ws["B2"].fill = dark_navy

# Tiles
tiles = [
    (2, "PORTFOLIO CM %", "='02_Unit_P&L'!H13", pct),
    (4, "ANN CONTRIB $000", "='02_Unit_P&L'!L13", money0),
    (6, "AFTER SLOTTING", "='02_Unit_P&L'!N13", money0),
    (9, "LAUNCH PAYBACK MO", "='04_Cohort_or_Ramp'!C21", "0"),
    (11, "Y1 NET $000", "='04_Cohort_or_Ramp'!C23", money0),
]
for sc, label, val, fmt in tiles:
    c = ws.cell(4, sc, label)
    c.font = tile_title
    c.fill = tile_fill
    c.border = thin
    c.alignment = Alignment(horizontal="center")
    if sc in (2, 4, 6, 9, 11):
        ws.merge_cells(start_row=4, start_column=sc, end_row=4, end_column=sc + 1)
        ws.cell(4, sc + 1).fill = tile_fill
        ws.cell(4, sc + 1).border = thin
    c = ws.cell(5, sc, val)
    c.font = tile_value
    c.border = thin
    c.number_format = fmt
    c.alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=5, start_column=sc, end_row=5, end_column=sc + 1)
    ws.cell(5, sc + 1).border = thin
ws.row_dimensions[5].height = 24

ws["B7"] = "SKU CONTRIBUTION RANK"
ws["B7"].font = section_font

for i, h in enumerate(["SKU", "CM %", "RAG", "Ann CM $000", "After slot $000"]):
    style_header_cell(ws.cell(8, 2 + i), h)

for i in range(len(SKUS)):
    r = 9 + i
    ws.cell(r, 2, f"='02_Unit_P&L'!B{6+i}")
    style_formula(ws.cell(r, 2))
    c = ws.cell(r, 3, f"='02_Unit_P&L'!H{6+i}")
    style_formula(c, key=True)
    c.number_format = pct
    c = ws.cell(r, 4, f"='02_Unit_P&L'!J{6+i}")
    style_formula(c, key=True)
    c = ws.cell(r, 5, f"='02_Unit_P&L'!L{6+i}")
    style_formula(c)
    c.number_format = money0
    c = ws.cell(r, 6, f"='02_Unit_P&L'!N{6+i}")
    style_formula(c, key=True)
    c.number_format = money0

ws.conditional_formatting.add("D9:D14", FormulaRule(formula=['D9="Green"'], fill=green_fill))
ws.conditional_formatting.add("D9:D14", FormulaRule(formula=['D9="Amber"'], fill=amber_fill))
ws.conditional_formatting.add("D9:D14", FormulaRule(formula=['D9="Red"'], fill=red_fill))

ws["B16"] = "CHANNEL TOTALS ($000s)"
ws["B16"].font = section_font
for i, h in enumerate(["Grocery", "Club", "eComm", "Convenience", "Total"]):
    style_header_cell(ws.cell(17, 2 + i), h)
for i in range(5):
    c = ws.cell(18, 2 + i, f"='03_Channel_View'!{get_column_letter(3+i)}27")
    style_formula(c, key=True)
    c.number_format = money0

ws["I7"] = "LAUNCH SNAPSHOT (NL-PROBIO)"
ws["I7"].font = section_font
ws.merge_cells("I7:L7")

launch_rows = [
    ("Unit CM $ after fee", "='04_Cohort_or_Ramp'!C9", money),
    ("FY volume (000)", "='04_Cohort_or_Ramp'!O13", num1),
    ("FY contribution $000", "='04_Cohort_or_Ramp'!O16", money0),
    ("Launch spend $000", "='04_Cohort_or_Ramp'!O14", money0),
    ("Y1 net $000", "='04_Cohort_or_Ramp'!C23", money0),
    ("Payback month", "='04_Cohort_or_Ramp'!C21", "0"),
    ("Within hurdle?", "='04_Cohort_or_Ramp'!C22", "@"),
]
for i, (label, ref, fmt) in enumerate(launch_rows):
    r = 8 + i
    ws.cell(r, 9, label).font = bold
    ws.cell(r, 9).border = thin
    c = ws.cell(r, 10, ref)
    style_formula(c, key=True)
    if fmt != "@":
        c.number_format = fmt
    ws.merge_cells(start_row=r, start_column=10, end_row=r, end_column=12)

ws["I16"] = "STRESS TOGGLES (from Assumptions)"
ws["I16"].font = section_font
ws.merge_cells("I16:L16")
ws["I17"] = "Price stress %"
ws["J17"] = "='01_Assumptions'!C18"
style_formula(ws["J17"])
ws["J17"].number_format = pct
ws["I18"] = "COGS stress %"
ws["J18"] = "='01_Assumptions'!C19"
style_formula(ws["J18"])
ws["J18"].number_format = pct
ws["I19"] = "Target CM %"
ws["J19"] = "='01_Assumptions'!C20"
style_formula(ws["J19"])
ws["J19"].number_format = pct

# Mini chart — SKU ann contrib
ws["B20"] = "Chart data"
ws["B20"].font = small_grey
ws["B21"] = "SKU"
ws["C21"] = "Ann CM $000"
for i in range(len(SKUS)):
    ws.cell(22 + i, 2, f"='02_Unit_P&L'!B{6+i}")
    c = ws.cell(22 + i, 3, f"='02_Unit_P&L'!L{6+i}")
    c.number_format = money0

chart = BarChart()
chart.type = "col"
chart.title = "Annual contribution by SKU ($000s)"
chart.style = 10
data = Reference(ws, min_col=3, min_row=21, max_row=27)
cats = Reference(ws, min_col=2, min_row=22, max_row=27)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.shape = 4
chart.width = 14
chart.height = 8
ws.add_chart(chart, "B29")

for col in range(2, 14):
    ws.cell(20, col).fill = dark_navy
# Don't darken chart data — put footer elsewhere
# Clear mistaken dark on row 20 — chart data label lives there
ws["B20"].fill = PatternFill()
ws["B20"].font = small_grey
for col in range(3, 14):
    ws.cell(20, col).fill = PatternFill()

for col in range(2, 14):
    ws.cell(28, col).fill = dark_navy
ws.merge_cells("B28:M28")
ws["B28"] = "Edit Assumptions (yellow) → Unit P&L / Channel / Ramp recalculate → this dashboard updates  ·  fictional sample"
ws["B28"].font = Font(name="Calibri", size=8, italic=True, color="D6E3F0")
ws["B28"].fill = dark_navy

ws.print_area = "A1:M40"
ws.page_margins.left = 0.4
ws.page_margins.right = 0.4
ws.page_margins.top = 0.4
ws.page_margins.bottom = 0.4

# ═══════════════════════════════════════════════════════════════════
# 06_Data_Dictionary
# ═══════════════════════════════════════════════════════════════════
ws = wb.create_sheet(DICT)
landscape(ws)
set_col_widths(ws, [3, 28, 18, 70])

ws["B2"] = "06  ·  Data Dictionary"
ws["B2"].font = title_font
ws["B3"] = "Field definitions so another analyst can inherit the file."
ws["B3"].font = italic_grey

for i, h in enumerate(["Field", "Tab", "Definition"]):
    style_header_cell(ws.cell(5, 2 + i), h)

defs = [
    ("SKU / Product", ASSUMP, "Fictional Northline assortment codes and names."),
    ("Price $", ASSUMP, "List / shelf price per unit before channel fees."),
    ("COGS $", ASSUMP, "Variable product cost per unit."),
    ("Var sell $", ASSUMP, "Variable selling cost per unit (freight-out, broker, etc.)."),
    ("Slotting $000", ASSUMP, "Annualized slotting / listing fees in $000s. Launch SKU also has monthly CAC on the ramp tab."),
    ("Ann vol (000)", ASSUMP, "Annual volume in thousands of units."),
    ("Price / COGS stress %", ASSUMP, "Global toggles applied on Unit P&L to sensitivity-test margin."),
    ("Target CM %", ASSUMP, "Hurdle contribution margin for RAG."),
    ("Unit contrib $", UNIT, "Price − COGS − Var sell (after stress). Key output."),
    ("CM %", UNIT, "Unit contrib ÷ Price."),
    ("Ann contrib $000", UNIT, "Unit contrib $ × annual volume (000)."),
    ("Ann CM after slot $000", UNIT, "Annual contribution minus slotting $000s."),
    ("Channel mix %", CHAN, "Share of SKU volume by Grocery / Club / eComm / Convenience. Yellow inputs."),
    ("Channel fee %", CHAN, "Retailer or platform take as % of price. Yellow inputs."),
    ("Unit CM $ after fee", CHAN, "Unit contrib − Price × blended fee %."),
    ("Channel $000s", CHAN, "Volume × mix × unit CM after fee, by channel."),
    ("Ramp volume / spend", RAMP, "Monthly 000 units and launch spend $000s for NL-PROBIO (yellow)."),
    ("Cumulative net $000s", RAMP, "Running sum of (monthly contribution − launch spend). Payback when it crosses ≥ 0."),
    ("Payback month #", RAMP, "First month cumulative net ≥ 0. Compared to hurdle on Assumptions."),
    ("Dashboard", DASH, "One-pager: portfolio tiles, SKU rank, channel totals, launch snapshot."),
    ("Color convention", "All", "Yellow + blue font = inputs. Black font = formulas. Green/Amber/Red = RAG."),
    ("Units", "All", "Per-unit in $. Annual flows in $000s. Volume in 000 units."),
]
for i, (field, tab, definition) in enumerate(defs):
    r = 6 + i
    ws.cell(r, 2, field).font = bold
    ws.cell(r, 2).border = thin
    ws.cell(r, 3, tab).border = thin
    cell = ws.cell(r, 4, definition)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    cell.border = thin
    ws.row_dimensions[r].height = 30

ws["B29"] = "All sample numbers are fictional. Built for a public GitHub portfolio — no employer data."
ws["B29"].font = small_grey

out = Path(__file__).resolve().parent / "Northline_Unit_Economics.xlsx"
wb.save(out)
print("Wrote", out)
print("Sheets:", wb.sheetnames)
